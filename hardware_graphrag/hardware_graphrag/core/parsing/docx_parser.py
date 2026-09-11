"""
DOCX parsing for hardware specification documents.

Uses python-docx to walk paragraphs (respecting Word's built-in heading
styles "Heading 1..N" when present, falling back to the same numbered-
heading regex heuristics used for PDFs) and tables.
"""

from __future__ import annotations

import re
from typing import List, Optional

from core.parsing.base_parser import BaseParser, HeadingClassifier
from utils.models import DocSection, ParsedDocument, new_id
from utils.logger import get_logger

logger = get_logger("parsing.docx")

try:
    import docx  # python-docx
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False

_WORD_HEADING_STYLE_RE = re.compile(r"Heading\s*(\d+)", re.IGNORECASE)


class DocxParser(BaseParser):
    file_type = "docx"

    def parse(self) -> ParsedDocument:
        doc_id = new_id("doc")
        if not _HAS_DOCX:
            logger.warning("python-docx not installed - cannot parse DOCX properly.")
            return ParsedDocument(doc_id=doc_id, filename=self.filename, file_type=self.file_type)

        document = docx.Document(self.filepath)
        sections: List[DocSection] = []
        current_chapter = "Untitled"
        current_section: Optional[DocSection] = None
        text_buf: List[str] = []
        raw_len = 0
        pseudo_page = 1  # DOCX has no fixed pagination without rendering; approximate by paragraph blocks

        def flush(sec, buf):
            if sec is not None and buf:
                sec.text = (sec.text + "\n" + "\n".join(buf)).strip()
            buf.clear()

        for para in document.paragraphs:
            line_text = para.text.strip()
            raw_len += len(line_text)
            if not line_text:
                continue

            style_name = para.style.name if para.style else ""
            style_match = _WORD_HEADING_STYLE_RE.match(style_name)
            heading_match = HeadingClassifier.match_heading(line_text)

            if style_match or heading_match:
                flush(current_section, text_buf)
                if style_match:
                    level = max(int(style_match.group(1)) - 1, 0)
                    number, title = "", line_text
                else:
                    number, level, title = heading_match

                sec = self._new_section(
                    title=title, level=min(level, 4), section_number=number,
                    page=pseudo_page, content_type="text",
                )
                if level == 0:
                    current_chapter = title
                sec.chapter = current_chapter
                sections.append(sec)
                current_section = sec
                pseudo_page += 1
                continue

            content_type = HeadingClassifier.classify_content_type(line_text)
            if content_type != "text":
                flush(current_section, text_buf)
                special = self._new_section(
                    title=line_text[:80], level=(current_section.level + 1) if current_section else 1,
                    page=pseudo_page, content_type=content_type, text=line_text,
                )
                special.chapter = current_chapter
                sections.append(special)
                continue

            if current_section is None:
                current_section = self._new_section(title=current_chapter, level=0, page=pseudo_page)
                current_section.chapter = current_chapter
                sections.append(current_section)
            text_buf.append(line_text)

        flush(current_section, text_buf)

        # tables
        for t_idx, table in enumerate(document.tables):
            rows_text = "\n".join(
                " | ".join(cell.text.strip() for cell in row.cells) for row in table.rows
            )
            if rows_text.strip():
                sec = self._new_section(
                    title=f"Table #{t_idx + 1}", level=2, page=pseudo_page,
                    content_type="table", text=rows_text,
                )
                sec.chapter = current_chapter
                sections.append(sec)

        self._link_sections(sections)
        logger.info(f"Parsed DOCX '{self.filename}': {len(sections)} sections.")
        return ParsedDocument(
            doc_id=doc_id, filename=self.filename, file_type=self.file_type,
            sections=sections, total_pages=pseudo_page, raw_text_len=raw_len,
        )
